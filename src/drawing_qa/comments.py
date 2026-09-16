"""Copy commented drawing PDFs into Sorted/{revision} using title-block detection."""

from __future__ import annotations

import shutil
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from drawing_qa.checker import iter_pdfs
from drawing_qa.config_loader import AppConfig
from drawing_qa.detect import extract_titleblock, recover_unlabelled_fields
from drawing_qa.extract import clear_page_word_cache, require_pymupdf
from drawing_qa.filename import parse_filename
from drawing_qa.models import TitleBlockFields
from drawing_qa.paths import next_available_report_path
from drawing_qa.tokens import format_pc_revision, parse_pc_revision

SORTED_ROOT = "Sorted"
UNSORTED_NAME = "Unsorted"


@dataclass
class CommentSortItem:
    source: Path
    revision: str | None = None
    folder_name: str = UNSORTED_NAME
    dest: Path | None = None
    title: str | None = None
    document_reference: str | None = None
    layout_id: str | None = None
    notes: str = ""
    copied: bool = False


@dataclass
class CommentSortResult:
    folder: Path
    sorted_root: Path
    items: list[CommentSortItem] = field(default_factory=list)
    report_path: Path | None = None

    def counts(self) -> Counter[str]:
        return Counter(item.folder_name for item in self.items)


def revision_folder_name(revision: str | None) -> str | None:
    """Return C01 / P02 for a title-block P/C revision, else None (manual sort)."""
    parsed = parse_pc_revision(revision)
    if not parsed:
        return None
    return format_pc_revision(*parsed)


def extract_comment_fields(
    path: Path,
    config: AppConfig,
    *,
    unsorted_name: str = UNSORTED_NAME,
) -> CommentSortItem:
    """Read the title-block revision with the same layouts as QA-TB-Checker."""
    filename = parse_filename(
        path,
        field_count=config.field_count,
        revision_pattern=config.revision_pattern,
    )
    item = CommentSortItem(
        source=path,
        title=filename.title,
        document_reference=filename.document_reference,
        folder_name=unsorted_name,
    )
    page = None
    try:
        require_pymupdf()
        import pymupdf

        doc = pymupdf.open(path)
        try:
            if doc.page_count < 1:
                item.notes = "PDF has no pages"
                return item
            page = doc[0]
            titleblock = extract_titleblock(
                page,
                config.layouts,
                config.min_layout_score,
            )
            recover_unlabelled_fields(page, config.layouts, titleblock)
        finally:
            if page is not None:
                clear_page_word_cache(page)
            doc.close()
    except Exception as exc:  # noqa: BLE001 - one bad PDF must not stop the folder
        clear_page_word_cache(page)
        item.notes = f"Could not read PDF: {exc}"
        return item

    item.layout_id = titleblock.layout_id
    item.title = titleblock.title or item.title
    item.document_reference = titleblock.document_reference or item.document_reference
    raw_rev = (titleblock.revision or "").strip()
    folder = revision_folder_name(raw_rev)
    if folder:
        item.revision = folder
        item.folder_name = folder
        return item
    if _layout_undetected(titleblock):
        item.notes = "Title block could not be read"
    elif not raw_rev:
        item.notes = "No revision in the title block"
    else:
        item.notes = f"Title-block revision {raw_rev!r} is not a P/C issue"
        item.revision = raw_rev
    return item


def _layout_undetected(titleblock: TitleBlockFields) -> bool:
    return not titleblock.layout_id or (
        titleblock.notes and any("below threshold" in note.lower() for note in titleblock.notes)
    )


def _copy_item(item: CommentSortItem, dest_dir: Path) -> None:
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / item.source.name
    shutil.copy2(item.source, dest)
    item.dest = dest
    item.copied = True


def organise_comment_pdfs(
    folder: Path,
    config: AppConfig,
    *,
    sorted_name: str = SORTED_ROOT,
    unsorted_name: str = UNSORTED_NAME,
    pdfs: list[Path] | None = None,
    on_pdf=None,
) -> CommentSortResult:
    """Copy PDFs into Sorted/{revision}; originals stay in the dump folder."""
    folder = folder.resolve()
    sorted_root = folder / sorted_name
    paths = pdfs if pdfs is not None else iter_pdfs(folder)
    skip_root = sorted_root.resolve()
    paths = [
        path
        for path in paths
        if skip_root not in path.resolve().parents and path.resolve() != skip_root
    ]
    result = CommentSortResult(folder=folder, sorted_root=sorted_root)
    total = len(paths)
    for index, path in enumerate(paths, start=1):
        item = extract_comment_fields(path, config, unsorted_name=unsorted_name)
        dest_dir = sorted_root / item.folder_name
        try:
            _copy_item(item, dest_dir)
        except OSError as exc:
            item.copied = False
            extra = f"Could not copy: {exc}"
            item.notes = f"{item.notes}; {extra}" if item.notes else extra
        result.items.append(item)
        if on_pdf is not None:
            on_pdf(index, total, item)
    return result


def default_comments_report_path(
    folder: Path, *, when: datetime | None = None
) -> Path:
    when = when or datetime.now()
    name = f"Comments_Organiser_{when.strftime('%d%m%y')}.xlsx"
    return next_available_report_path(folder, name)


def write_comments_report(result: CommentSortResult, path: Path | None = None) -> Path:
    path = path or default_comments_report_path(result.folder)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _write_summary(summary, result)
    listing = wb.create_sheet("Comments")
    _write_comments(listing, result)
    wb.save(path)
    result.report_path = path
    return path


_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
_TITLE_FONT = Font(bold=True, size=14, color="1F4E79", name="Calibri")
_BODY = Font(name="Calibri", size=11)
_INTRO = Font(name="Calibri", size=11, italic=True)
_UNSORTED_FILL = PatternFill("solid", fgColor="FFEB9C")
_OK_FILL = PatternFill("solid", fgColor="C6EFCE")
_MISSING_FILL = PatternFill("solid", fgColor="FFC7CE")


def _sorted_items(result: CommentSortResult) -> list[CommentSortItem]:
    return sorted(
        result.items,
        key=lambda item: (
            _revision_sort_key(item.folder_name),
            (item.document_reference or "").upper(),
            item.source.name.lower(),
        ),
    )


def _revision_folders(result: CommentSortResult) -> list[str]:
    counts = result.counts()
    folders = sorted(
        (name for name in counts if name != UNSORTED_NAME),
        key=_revision_sort_key,
    )
    if UNSORTED_NAME in counts:
        folders.append(UNSORTED_NAME)
    return folders


def _write_summary(ws, result: CommentSortResult) -> None:
    when = datetime.now().strftime("%d/%m/%y")
    ws["A1"] = "Drawing comments"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:B1")
    ws["A2"] = (
        "Commented PDFs have been grouped by the revision printed on the drawing. "
        "Please action the comments against that issue."
    )
    ws["A2"].font = _INTRO
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.merge_cells("A2:B2")
    ws.row_dimensions[2].height = 36
    ws["A3"] = "Date"
    ws["B3"] = when
    ws["A3"].font = _BODY
    ws["B3"].font = _BODY

    ws["A5"] = "Revision"
    ws["B5"] = "Comments files"
    for cell in (ws["A5"], ws["B5"]):
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    counts = result.counts()
    row = 6
    for name in _revision_folders(result):
        label = "Not identified" if name == UNSORTED_NAME else name
        ws.cell(row, 1, label)
        ws.cell(row, 2, counts[name])
        fill = _UNSORTED_FILL if name == UNSORTED_NAME else _OK_FILL
        ws.cell(row, 1).fill = fill
        for col in (1, 2):
            ws.cell(row, col).font = _BODY
        row += 1
    ws.cell(row + 1, 1, "Total")
    ws.cell(row + 1, 2, len(result.items))
    ws.cell(row + 1, 1).font = Font(name="Calibri", size=11, bold=True)
    ws.cell(row + 1, 2).font = Font(name="Calibri", size=11, bold=True)
    if UNSORTED_NAME in counts:
        ws.cell(row + 3, 1, "Not identified: revision could not be read. Sort these by hand.")
        ws.cell(row + 3, 1).font = _INTRO
        ws.merge_cells(start_row=row + 3, start_column=1, end_row=row + 3, end_column=2)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 18


def _write_comments(ws, result: CommentSortResult) -> None:
    headers = [
        "Revision",
        "Document reference",
        "Title",
        "File name",
        "Comments file",
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(1, col, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for row, item in enumerate(_sorted_items(result), start=2):
        exists = bool(item.copied and item.dest is not None and item.dest.is_file())
        revision = (
            "Not identified"
            if item.folder_name == UNSORTED_NAME
            else (item.revision or item.folder_name)
        )
        values = [
            revision,
            item.document_reference or "",
            item.title or "",
            item.source.name,
            "Yes" if exists else "No",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            cell.font = _BODY
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        if item.folder_name == UNSORTED_NAME:
            ws.cell(row, 1).fill = _UNSORTED_FILL
        else:
            ws.cell(row, 1).fill = _OK_FILL
        ws.cell(row, 5).fill = _OK_FILL if exists else _MISSING_FILL
        ws.row_dimensions[row].height = 22
    widths = (16, 38, 48, 52, 16)
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width
    last = max(1, len(result.items) + 1)
    ws.auto_filter.ref = f"A1:E{last}"
    ws.freeze_panes = "A2"


def _revision_sort_key(name: str) -> tuple:
    parsed = parse_pc_revision(name)
    if parsed:
        series, number = parsed
        return (0 if series == "P" else 1, number, name)
    return (9, 0, name)

