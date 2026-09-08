from pathlib import Path

from openpyxl import Workbook

from drawing_qa.checker import check_paths
from drawing_qa.config_loader import load_config
from drawing_qa.document_list import (
    DocumentListIndex,
    PortalDocument,
    check_document_list,
    find_document_list,
    load_document_list,
)
from drawing_qa.models import (
    CheckStatus,
    DocumentResult,
    FilenameFields,
    TitleBlockFields,
    finalize_status,
)
from drawing_qa.paths import bundled_config_dir
from drawing_qa.tokens import is_allowed_first_revision, is_successor_revision, next_revision
from tests.pdf_fixtures import write_bottom_right_pdf


def _layout():
    return load_config(bundled_config_dir()).document_list.layout


def _write_excel(
    path: Path,
    rows: list[tuple[str, str, str]],
    *,
    headers: list[str] | None = None,
    header_row: int = 1,
) -> Path:
    wb = Workbook()
    ws = wb.active
    for _ in range(header_row - 1):
        ws.append(["ignore"] * 3)
    ws.append(headers or ["Original Doc Ref (Non-Standard)", "Description", "Revision"])
    for row in rows:
        ws.append(list(row))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _drawing(
    *,
    project: str = "ABC",
    number: str = "0001",
    revision: str = "P01",
    title: str = "Ground Floor GA",
    doc_ref: str | None = None,
    suitability: str | None = None,
) -> DocumentResult:
    ref = doc_ref or f"{project}-WXY-ZZ-00-DR-A-{number}"
    return DocumentResult(
        path=Path(f"{ref}-{revision}.pdf"),
        filename=FilenameFields(
            raw_stem=f"{ref}-{revision}",
            document_reference=ref,
            title=title,
            revision=revision,
            parse_ok=True,
            parts={"project": project},
        ),
        titleblock=TitleBlockFields(
            document_reference=ref,
            title=title,
            revision=revision,
            suitability=suitability,
        ),
        status=CheckStatus.MATCH,
    )
    ref = doc_ref or f"{project}-WXY-ZZ-00-DR-A-{number}"
    return DocumentResult(
        path=Path(f"{ref}-{revision}.pdf"),
        filename=FilenameFields(
            raw_stem=f"{ref}-{revision}",
            document_reference=ref,
            title=title,
            revision=revision,
            parse_ok=True,
            parts={"project": project},
        ),
        titleblock=TitleBlockFields(
            document_reference=ref,
            title=title,
            revision=revision,
        ),
        status=CheckStatus.MATCH,
    )


def test_successor_revision_rules():
    assert is_successor_revision("P01", "P02")
    assert is_successor_revision("P01", "C01")
    assert is_successor_revision("P09", "C01")
    assert not is_successor_revision("P01", "P01")
    assert not is_successor_revision("P01", "P03")
    assert not is_successor_revision("P01", "C02")
    assert not is_successor_revision("C01", "P02")
    assert is_successor_revision(".C04", "C05")
    assert is_successor_revision("C01.", "C02")
    assert next_revision("P01") == "P02"
    assert next_revision(".C04") == "C05"
    assert next_revision("C01.") == "C02"
    assert is_allowed_first_revision("C01", ["P01", "C01"])
    assert not is_allowed_first_revision("P02", ["P01", "C01"])


def test_load_4projects_headers(tmp_path: Path):
    path = _write_excel(
        tmp_path / "OVCD Document Listing.xlsx",
        [("R459-MBS-DZ-ZZ-DR-W-0001", "Plant", "P01")],
    )
    index = load_document_list(path, _layout())
    row = index.get("R459-MBS-DZ-ZZ-DR-W-0001")
    assert row is not None
    assert row.revision == "P01"
    assert row.title == "Plant"


def test_oval_prefers_original_doc_ref_when_name_also_has_a_number(tmp_path: Path):
    """Oval C+D 4Projects: Name can look like a doc ref; the real number is Original Doc Ref."""
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Name",
            "Description",
            "Revision",
            "Original Doc Ref (Non-Standard)",
            "Revision Workflow",
        ]
    )
    ws.append(
        [
            "R459-MBS-DZ-ZZ-DR-W-99999",
            "Plant",
            "P01",
            "R459-MBS-DZ-ZZ-DR-W-0001",
            "Under Review",
        ]
    )
    path = tmp_path / "OVCD Document Listing.xlsx"
    wb.save(path)
    index = load_document_list(path, _layout())
    assert index.get("R459-MBS-DZ-ZZ-DR-W-0001") is not None
    assert index.get("R459-MBS-DZ-ZZ-DR-W-99999") is None


def _write_4projects_name_listing(
    path: Path, doc_ref: str, title: str, revision: str = "P01"
) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Items"
    ws.append(
        [
            "Name",
            "Description",
            "Revision",
            "Original Doc Ref (Non-Standard)",
            "Revision Workflow",
            "Status",
        ]
    )
    ws.append([doc_ref, title, revision, "", "Under Review", "S3 - For Review & Comment"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def test_uses_name_when_original_doc_ref_column_is_blank(tmp_path: Path):
    """Trillium 4Projects dumps leave Original Doc Ref empty; ISO numbers are in Name."""
    path = _write_4projects_name_listing(
        tmp_path / "Tril Document Listing 080926.xlsx",
        "R456-MAL20-BI-ZZ-DR-W-605-001",
        "Block I - Combined Services - Apt Type I.01a,b,c,d",
    )
    index = load_document_list(path, _layout())
    row = index.get("R456-MAL20-BI-ZZ-DR-W-605-001")
    assert row is not None
    assert row.revision == "P01"
    assert row.title.startswith("Block I - Combined Services")
    assert row.status == "Under Review"

    result = check_document_list(
        [
            _drawing(
                project="R456",
                number="605-001",
                revision="P01",
                title="Block I - Combined Services - Apt Type I.01a,b,c,d",
                doc_ref="R456-MAL20-BI-ZZ-DR-W-605-001",
            )
        ],
        index,
        _layout(),
    )[0]
    finalize_status(result)
    assert result.portal_revision == "P01"
    assert CheckStatus.PORTAL_REVISION in result.issues


def test_wcr_uses_name_column_like_trillium(tmp_path: Path):
    path = _write_4projects_name_listing(
        tmp_path / "WCR Document Listing.xlsx",
        "WCR-MBS-B7-ZZ-DR-E-6105",
        "Apartment Type B7-2C DUP",
        "C04",
    )
    index = load_document_list(path, _layout())
    row = index.get("WCR-MBS-B7-ZZ-DR-E-6105")
    assert row is not None
    assert row.revision == "C04"


def test_name_column_mapping_follows_pdf_project_when_listing_name_is_generic(
    tmp_path: Path,
):
    path = _write_4projects_name_listing(
        tmp_path / "Document Listing.xlsx",
        "R456-MAL20-BI-ZZ-DR-W-605-001",
        "Combined Services",
    )
    index = load_document_list(path, _layout(), project_codes=["R456"])
    assert index.get("R456-MAL20-BI-ZZ-DR-W-605-001") is not None


def test_strips_leading_and_trailing_dots_from_portal_revision(tmp_path: Path):
    path = _write_excel(
        tmp_path / "WCR Listing.xlsx",
        [("WCR-MBS-B7-ZZ-DR-E-6105", "Apartment Type B7-2C DUP", ".C04")],
    )
    index = load_document_list(path, _layout())
    row = index.get("WCR-MBS-B7-ZZ-DR-E-6105")
    assert row is not None
    assert row.revision == "C04"

    layout = _layout()
    result = check_document_list(
        [_drawing(project="WCR", number="6105", revision="C05", doc_ref="WCR-MBS-B7-ZZ-DR-E-6105")],
        index,
        layout,
    )[0]
    finalize_status(result)
    assert result.portal_revision == "C04"
    assert CheckStatus.PORTAL_REVISION not in result.issues


def test_load_asite_headers_on_row_six(tmp_path: Path):
    path = _write_excel(
        tmp_path / "Asite export.xlsx",
        [("J106309-MBS-ZZ-00-DR-M-0001", "Basement", "P03")],
        headers=["Doc Ref", "Doc Title", "Rev"],
        header_row=6,
    )
    index = load_document_list(path, _layout())
    row = index.get("J106309-MBS-ZZ-00-DR-M-0001")
    assert row is not None
    assert row.revision == "P03"
    assert row.title == "Basement"


def test_load_dochosting_csv(tmp_path: Path):
    path = tmp_path / "HPA dump.csv"
    path.write_text(
        "Title,Subject,Rev\nHPA-MBS-ZZ-00-DR-E-0001,Lighting,P01\n",
        encoding="utf-8",
    )
    index = load_document_list(path, _layout())
    row = index.get("HPA-MBS-ZZ-00-DR-E-0001")
    assert row is not None
    assert row.revision == "P01"
    assert row.title == "Lighting"


def test_holloway_uses_title_and_description_not_subject(tmp_path: Path):
    """DocHosting HP dump: Title is the ISO number, Description is the drawing title."""
    path = tmp_path / "HP Document Listing 070926.csv"
    path.write_text(
        "Report Created,Project Folder,Title,Subject,Description,Status,Rev,Date\n"
        "06-09-2026,/ACAD_Services,HPA-MBS-C1-ZZ-SM-X-52007,Schematics,"
        "Block C1 LTHW M-BUS schematic,Construction,C1,13-Aug-25\n",
        encoding="utf-8",
    )
    index = load_document_list(path, _layout())
    row = index.get("HPA-MBS-C1-ZZ-SM-X-52007")
    assert row is not None
    assert row.revision == "C1"
    assert row.title == "Block C1 LTHW M-BUS schematic"
    assert row.status == "Construction"
    from drawing_qa.document_list import status_allows_upload

    assert status_allows_upload(row.status, _layout(), "HPA")


def test_barking_uses_asite_status_not_workflow_status(tmp_path: Path):
    """Asite listings have Status (For Status Change) and Workflow Status (RUNNING)."""
    wb = Workbook()
    ws = wb.active
    for _ in range(6):
        ws.append(["ignore"])
    ws.append(["Status", "Type", "Doc Ref", "Doc Title", "Rev", "Workflow Status"])
    ws.append(
        [
            "For Status Change",
            "pdf",
            "J106309-MBS-ZZ-ZZ-DR-X-580010",
            "26SSD PT8 STRUCTURAL OPENING",
            "P01",
            "RUNNING",
        ]
    )
    ws.append(
        [
            "B - Partial Sign Off (with comment)",
            "pdf",
            "J106309-MBS-ZZ-ZZ-DR-X-620301",
            "EXTERNAL COMBINED SERVICES LAYOUT",
            "P05",
            "COMPLETED",
        ]
    )
    path = tmp_path / "BR Document Listing 070926.xlsx"
    wb.save(path)
    index = load_document_list(path, _layout())
    blocked = index.get("J106309-MBS-ZZ-ZZ-DR-X-580010")
    allowed = index.get("J106309-MBS-ZZ-ZZ-DR-X-620301")
    assert blocked is not None
    assert blocked.revision == "P01"
    assert blocked.title == "26SSD PT8 STRUCTURAL OPENING"
    assert blocked.status == "For Status Change"
    assert allowed is not None
    assert allowed.status == "B - Partial Sign Off (with comment)"
    from drawing_qa.document_list import status_allows_upload

    layout = _layout()
    assert not status_allows_upload(blocked.status, layout, "J106309")
    assert status_allows_upload(allowed.status, layout, "J106309")


def test_skips_asite_comments_report(tmp_path: Path):
    layout = _layout()
    comments = tmp_path / "BR Comments Report 070926.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Comments Report"])
    ws.append([])
    ws.append(["Doc Ref", "Document Title", "Rev"])
    ws.append(["J106309-MBS-ZZ-00-DR-M-0001", "Basement", "P99"])
    wb.save(comments)
    listing = _write_excel(
        tmp_path / "BR Document Listing 070926.xlsx",
        [("J106309-MBS-ZZ-00-DR-M-0001", "Basement", "P03")],
        headers=["Doc Ref", "Doc Title", "Rev"],
    )
    assert find_document_list(tmp_path, layout) == listing
    only_comments = tmp_path / "comments_only"
    only_comments.mkdir()
    comments.replace(only_comments / comments.name)
    assert find_document_list(only_comments, layout) is None


def test_keeps_highest_portal_revision(tmp_path: Path):
    path = _write_excel(
        tmp_path / "Listing.xlsx",
        [
            ("ABC-WXY-ZZ-00-DR-A-0001", "Ground Floor GA", "P01"),
            ("ABC-WXY-ZZ-00-DR-A-0001", "Ground Floor GA", "P02"),
        ],
    )
    index = load_document_list(path, _layout())
    assert index.get("ABC-WXY-ZZ-00-DR-A-0001").revision == "P02"


def test_skips_irs_and_tbcheck_report_names(tmp_path: Path):
    layout = _layout()
    irs = _write_excel(
        tmp_path / "OVCD IRS.xlsx",
        [("ABC-WXY-ZZ-00-DR-A-0001", "Ground Floor GA", "P01")],
    )
    report = _write_excel(
        tmp_path / "ABC_260826.xlsx",
        [("ABC-WXY-ZZ-00-DR-A-0001", "Ground Floor GA", "P01")],
    )
    listing = _write_excel(
        tmp_path / "Document Listing.xlsx",
        [("ABC-WXY-ZZ-00-DR-A-0001", "Ground Floor GA", "P02")],
    )
    assert find_document_list(tmp_path, layout) == listing
    assert find_document_list(tmp_path, layout, explicit=irs) == irs
    only_irs = tmp_path / "irs_only"
    only_irs.mkdir()
    _write_excel(
        only_irs / "Project IRS.xlsx",
        [("ABC-WXY-ZZ-00-DR-A-0001", "Ground Floor GA", "P01")],
    )
    assert find_document_list(only_irs, layout) is None


def test_portal_successor_and_same_revision():
    layout = _layout()
    index = DocumentListIndex(
        path=Path("Document Listing.xlsx"),
        by_ref={
            "ABC-WXY-ZZ-00-DR-A-0001": PortalDocument(
                "ABC-WXY-ZZ-00-DR-A-0001", "P01", "Ground Floor GA"
            )
        },
    )
    ok = check_document_list([_drawing(revision="P02")], index, layout)[0]
    finalize_status(ok)
    assert CheckStatus.PORTAL_REVISION not in ok.issues

    same = check_document_list([_drawing(revision="P01")], index, layout)[0]
    finalize_status(same)
    assert CheckStatus.PORTAL_REVISION in same.issues

    skip = check_document_list([_drawing(revision="P03")], index, layout)[0]
    finalize_status(skip)
    assert CheckStatus.PORTAL_REVISION in skip.issues

    construction = check_document_list([_drawing(revision="C01")], index, layout)[0]
    finalize_status(construction)
    assert CheckStatus.PORTAL_REVISION not in construction.issues


def test_wcr_allows_c01_when_not_on_portal():
    layout = _layout()
    index = DocumentListIndex(path=Path("WCR Listing.xlsx"), by_ref={})
    first = check_document_list(
        [_drawing(project="WCR", revision="C01")], index, layout
    )[0]
    finalize_status(first)
    assert CheckStatus.PORTAL_REVISION not in first.issues

    too_far = check_document_list(
        [_drawing(project="WCR", revision="P02")], index, layout
    )[0]
    finalize_status(too_far)
    assert CheckStatus.PORTAL_REVISION in too_far.issues


def test_new_drawing_must_be_p01_except_wcr():
    layout = _layout()
    index = DocumentListIndex(path=Path("Listing.xlsx"), by_ref={})
    ok = check_document_list([_drawing(revision="P01")], index, layout)[0]
    finalize_status(ok)
    assert CheckStatus.PORTAL_REVISION not in ok.issues

    bad = check_document_list([_drawing(revision="C01")], index, layout)[0]
    finalize_status(bad)
    assert CheckStatus.PORTAL_REVISION in bad.issues


def test_portal_title_mismatch():
    layout = _layout()
    index = DocumentListIndex(
        path=Path("Listing.xlsx"),
        by_ref={
            "ABC-WXY-ZZ-00-DR-A-0001": PortalDocument(
                "ABC-WXY-ZZ-00-DR-A-0001", "P01", "Ground Floor GA"
            )
        },
    )
    result = check_document_list(
        [_drawing(revision="P02", title="Roof Plan")], index, layout
    )[0]
    finalize_status(result)
    assert CheckStatus.PORTAL_TITLE in result.issues


def test_check_paths_uses_folder_listing(tmp_path: Path, config_dir: Path):
    write_bottom_right_pdf(
        tmp_path / "ABC-WXY-ZZ-00-DR-A-0001-P01.pdf",
        document_reference="ABC-WXY-ZZ-00-DR-A-0001",
        title="Ground Floor GA",
        revision="P01",
    )
    _write_excel(
        tmp_path / "Document Listing.xlsx",
        [("ABC-WXY-ZZ-00-DR-A-0001", "Ground Floor GA", "P01")],
    )
    results = check_paths(
        list(tmp_path.glob("*.pdf")),
        load_config(config_dir),
    )
    assert results[0].portal_list_name == "Document Listing.xlsx"
    assert CheckStatus.PORTAL_REVISION in results[0].issues


def test_check_paths_skips_when_no_listing(tmp_path: Path, config_dir: Path):
    write_bottom_right_pdf(
        tmp_path / "ABC-WXY-ZZ-00-DR-A-0001-P01.pdf",
        document_reference="ABC-WXY-ZZ-00-DR-A-0001",
        title="Ground Floor GA",
        revision="P01",
    )
    results = check_paths(
        list(tmp_path.glob("*.pdf")),
        load_config(config_dir),
    )
    assert results[0].portal_list_name == ""
    assert CheckStatus.PORTAL_REVISION not in results[0].issues
    assert results[0].status == CheckStatus.MATCH


def test_status_allows_upload_uses_project_wordings():
    from drawing_qa.document_list import status_allows_upload

    layout = _layout()
    assert status_allows_upload("A Proceed", layout, "R459")
    assert status_allows_upload("Status A", layout, "R459")
    assert status_allows_upload("C", layout, "WCR")
    assert status_allows_upload("EA+DM - Status B", layout, "WCR")
    assert status_allows_upload("Construction", layout, "HPA")
    assert not status_allows_upload("Construction", layout, "R459")
    assert status_allows_upload("QA Approved", layout, "R459")
    assert not status_allows_upload("Pending QA Check", layout, "R459")
    assert not status_allows_upload("QA Rejected", layout, "R459")
    assert not status_allows_upload("", layout, "R459")


def test_prefers_revision_workflow_over_purpose_status_column(tmp_path: Path):
    path = _write_excel(
        tmp_path / "OVCD Document Listing.xlsx",
        [
            (
                "R459-MBS-DZ-ZZ-DR-W-0001",
                "Plant",
                "P01",
                "S3 - For Review & Comment",
                "Pending QA Check",
            )
        ],
        headers=[
            "Original Doc Ref (Non-Standard)",
            "Description",
            "Revision",
            "Status",
            "Revision Workflow",
        ],
    )
    index = load_document_list(path, _layout())
    assert index.has_status
    row = index.get("R459-MBS-DZ-ZZ-DR-W-0001")
    assert row is not None
    assert row.status == "Pending QA Check"


def test_blocks_upload_when_portal_status_is_not_abc():
    layout = _layout()
    index = DocumentListIndex(
        path=Path("Document Listing.xlsx"),
        has_status=True,
        by_ref={
            "ABC-WXY-ZZ-00-DR-A-0001": PortalDocument(
                "ABC-WXY-ZZ-00-DR-A-0001",
                "P01",
                "Ground Floor GA",
                "Pending QA Check",
            )
        },
    )
    blocked = check_document_list([_drawing(revision="P02")], index, layout)[0]
    assert blocked.portal_blocks_upload
    assert blocked.portal_status == "Pending QA Check"

    index.by_ref["ABC-WXY-ZZ-00-DR-A-0001"].status = "A Proceed"
    ok = check_document_list([_drawing(revision="P02")], index, layout)[0]
    assert not ok.portal_blocks_upload

    index.by_ref["ABC-WXY-ZZ-00-DR-A-0001"].status = "QA Approved"
    qa_ok = check_document_list([_drawing(revision="P02")], index, layout)[0]
    assert not qa_ok.portal_blocks_upload


def test_new_portal_drawing_is_not_an_upload_block():
    layout = _layout()
    index = DocumentListIndex(path=Path("Listing.xlsx"), has_status=True, by_ref={})
    result = check_document_list([_drawing(revision="P01")], index, layout)[0]
    assert not result.portal_blocks_upload


def test_intended_upload_revision_ignores_skipped_drawing_rev():
    from drawing_qa.document_list import intended_upload_revision

    assert intended_upload_revision("C01", "C02") == "C02"
    assert intended_upload_revision("C01", "C03") == "C02"
    assert intended_upload_revision("C01", "C01") == "C02"
    assert intended_upload_revision("P01", "C01") == "C01"
    assert intended_upload_revision("P01", "P02") == "P02"
    assert intended_upload_revision("P01", "P01") == "P02"


def test_check_sets_proposed_upload_to_next_portal_issue():
    layout = _layout()
    index = DocumentListIndex(
        path=Path("Document Listing.xlsx"),
        has_status=True,
        by_ref={
            "ABC-WXY-ZZ-00-DR-A-0001": PortalDocument(
                "ABC-WXY-ZZ-00-DR-A-0001",
                "C01",
                "Ground Floor GA",
                "Pending QA Check",
            )
        },
    )
    result = check_document_list([_drawing(revision="C03")], index, layout)[0]
    assert result.proposed_upload_revision == "C02"
    assert result.portal_blocks_upload


def test_r459_approved_p_must_go_to_c01_construction():
    layout = _layout()
    index = DocumentListIndex(
        path=Path("OVCD Document Listing.xlsx"),
        has_status=True,
        by_ref={
            "R459-WXY-ZZ-00-DR-A-0001": PortalDocument(
                "R459-WXY-ZZ-00-DR-A-0001",
                "P04",
                "Ground Floor GA",
                "A Proceed",
            )
        },
    )
    still_p = check_document_list(
        [_drawing(project="R459", revision="P05")], index, layout
    )[0]
    finalize_status(still_p)
    assert still_p.construction_upgrade_required
    assert still_p.proposed_upload_revision == "C01"
    assert CheckStatus.PORTAL_REVISION in still_p.issues
    assert CheckStatus.PURPOSE_MISMATCH not in still_p.issues

    ok = check_document_list(
        [_drawing(project="R459", revision="C01")],
        index,
        layout,
    )[0]
    finalize_status(ok)
    assert ok.proposed_upload_revision == "C01"
    assert CheckStatus.PORTAL_REVISION not in ok.issues

    # Purpose of issue is not pinned to one whitelist string here.
    review_purpose = check_document_list(
        [
            _drawing(
                project="R459",
                revision="C01",
                suitability="S3 - For Review & Comment",
            )
        ],
        index,
        layout,
    )[0]
    finalize_status(review_purpose)
    assert CheckStatus.PORTAL_REVISION not in review_purpose.issues
    assert CheckStatus.PURPOSE_MISMATCH not in review_purpose.issues


def test_r459_rejected_p_can_still_go_to_p_next():
    layout = _layout()
    index = DocumentListIndex(
        path=Path("OVCD Document Listing.xlsx"),
        has_status=True,
        by_ref={
            "R459-WXY-ZZ-00-DR-A-0001": PortalDocument(
                "R459-WXY-ZZ-00-DR-A-0001",
                "P04",
                "Ground Floor GA",
                "C Rejected",
            )
        },
    )
    result = check_document_list(
        [_drawing(project="R459", revision="P05")], index, layout
    )[0]
    finalize_status(result)
    assert not result.construction_upgrade_required
    assert result.proposed_upload_revision == "P05"
    assert CheckStatus.PORTAL_REVISION not in result.issues


def test_construction_upgrade_is_off_for_other_projects():
    layout = _layout()
    index = DocumentListIndex(
        path=Path("Listing.xlsx"),
        has_status=True,
        by_ref={
            "ABC-WXY-ZZ-00-DR-A-0001": PortalDocument(
                "ABC-WXY-ZZ-00-DR-A-0001",
                "P01",
                "Ground Floor GA",
                "A Proceed",
            )
        },
    )
    result = check_document_list([_drawing(revision="P02")], index, layout)[0]
    finalize_status(result)
    assert not result.construction_upgrade_required
    assert CheckStatus.PORTAL_REVISION not in result.issues


def test_intended_upload_revision_can_require_c01():
    from drawing_qa.document_list import intended_upload_revision

    assert intended_upload_revision("P04", "P05", require_revision="C01") == "C01"
