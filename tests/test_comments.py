from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from drawing_qa.cli_comments import main as comments_main
from drawing_qa.comments import (
    UNSORTED_NAME,
    organise_comment_pdfs,
    revision_folder_name,
    write_comments_report,
)
from drawing_qa.config_loader import load_config
from tests.pdf_fixtures import write_mbs_right_pdf, write_plain_pdf


def test_revision_folder_name_pads_pc_issues():
    assert revision_folder_name("C1") == "C01"
    assert revision_folder_name("p2") == "P02"
    assert revision_folder_name("C04") == "C04"
    assert revision_folder_name("S5") is None
    assert revision_folder_name("CONSTRUCTION") is None
    assert revision_folder_name(None) is None


def test_copies_mixed_revisions_and_leaves_originals(tmp_path: Path, config_dir: Path):
    dump = tmp_path / "comments"
    c02 = write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-10001.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-10001",
        title="Ground Floor",
        revision="C02",
    )
    c04 = write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-10002.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-10002",
        title="First Floor",
        revision="C04",
    )
    result = organise_comment_pdfs(dump, load_config(config_dir))
    folders = {item.source.name: item.folder_name for item in result.items}
    assert folders[c02.name] == "C02"
    assert folders[c04.name] == "C04"
    assert c02.is_file()
    assert c04.is_file()
    assert (dump / "Sorted" / "C02" / c02.name).is_file()
    assert (dump / "Sorted" / "C04" / c04.name).is_file()
    assert c02.read_bytes() == (dump / "Sorted" / "C02" / c02.name).read_bytes()


def test_pads_c1_folder_and_sends_undetected_to_unsorted(
    tmp_path: Path, config_dir: Path
):
    dump = tmp_path / "comments"
    c1 = write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-20001.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-20001",
        title="Roof",
        revision="C1",
    )
    plain = write_plain_pdf(dump / "mystery-markup.pdf")
    result = organise_comment_pdfs(dump, load_config(config_dir))
    by_name = {item.source.name: item for item in result.items}
    assert by_name[c1.name].folder_name == "C01"
    assert by_name[plain.name].folder_name == UNSORTED_NAME
    assert "title block" in by_name[plain.name].notes.lower()
    assert (dump / "Sorted" / "C01" / c1.name).is_file()
    assert (dump / "Sorted" / UNSORTED_NAME / plain.name).is_file()
    assert plain.is_file()


def test_does_not_rescan_copies_already_in_sorted(tmp_path: Path, config_dir: Path):
    dump = tmp_path / "comments"
    original = write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-30001.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-30001",
        title="Core",
        revision="P03",
    )
    decoy = write_plain_pdf(dump / "Sorted" / "C99" / "already-sorted.pdf")
    result = organise_comment_pdfs(dump, load_config(config_dir))
    names = [item.source.name for item in result.items]
    assert original.name in names
    assert decoy.name not in names
    assert decoy.is_file()
    assert not (dump / "Sorted" / UNSORTED_NAME / decoy.name).exists()


def test_writes_comments_report(tmp_path: Path, config_dir: Path):
    dump = tmp_path / "comments"
    write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-40004.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-40004",
        title="Roof",
        revision="C04",
    )
    write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-40002.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-40002",
        title="Plant",
        revision="P02",
    )
    write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-40003.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-40003",
        title="Stairs",
        revision="C02",
    )
    write_plain_pdf(dump / "no-block.pdf")
    result = organise_comment_pdfs(dump, load_config(config_dir))
    report = write_comments_report(result, dump / "out.xlsx")
    wb = load_workbook(report)
    summary = wb["Summary"]
    assert summary["A1"].value == "Drawing comments"
    assert summary["A5"].value == "Revision"
    assert summary["B5"].value == "Comments files"
    counts = {
        summary.cell(row, 1).value: summary.cell(row, 2).value for row in range(6, 10)
    }
    assert counts["P02"] == 1
    assert counts["C02"] == 1
    assert counts["C04"] == 1
    assert counts["Not identified"] == 1
    assert [summary.cell(row, 1).value for row in range(6, 10)] == [
        "P02",
        "C02",
        "C04",
        "Not identified",
    ]
    listing = wb["Comments"]
    headers = [listing.cell(1, col).value for col in range(1, 6)]
    assert headers == [
        "Revision",
        "Document reference",
        "Title",
        "File name",
        "Comments file",
    ]
    assert "Layout" not in headers
    rows = [
        [listing.cell(row, col).value for col in range(1, 6)] for row in range(2, 6)
    ]
    assert [row[0] for row in rows] == ["P02", "C02", "C04", "Not identified"]
    assert rows[0][1] == "R459-MBS-ZZ-ZZ-DR-W-40002"
    assert rows[0][2] == "Plant"
    assert rows[0][3] == "R459-MBS-ZZ-ZZ-DR-W-40002.pdf"
    assert rows[0][4] == "Yes"
    assert rows[3][3] == "no-block.pdf"
    assert rows[3][4] == "Yes"


def test_cli_comments_copies_and_reports(tmp_path: Path, config_dir: Path):
    dump = tmp_path / "comments"
    write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-50001.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-50001",
        title="Stairs",
        revision="C03",
    )
    output = dump / "report.xlsx"
    code = comments_main(
        [
            str(dump),
            "--config-dir",
            str(config_dir),
            "--output",
            str(output),
            "--no-pause",
        ]
    )
    assert code == 0
    assert output.is_file()
    assert (dump / "Sorted" / "C03" / "R459-MBS-ZZ-ZZ-DR-W-50001.pdf").is_file()


def test_cli_comments_dropped_pdfs_only(tmp_path: Path, config_dir: Path):
    dump = tmp_path / "comments"
    keep = write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-60001.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-60001",
        title="Kept",
        revision="C01",
    )
    skip = write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-60002.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-60002",
        title="Skipped",
        revision="C02",
    )
    code = comments_main(
        [
            str(keep),
            "--config-dir",
            str(config_dir),
            "--output",
            str(dump / "dropped.xlsx"),
            "--no-pause",
        ]
    )
    assert code == 0
    assert (dump / "Sorted" / "C01" / keep.name).is_file()
    assert not (dump / "Sorted" / "C02" / skip.name).exists()
    assert skip.is_file()


def test_cli_comments_exit_one_when_unsorted(tmp_path: Path, config_dir: Path):
    dump = tmp_path / "comments"
    write_plain_pdf(dump / "unknown.pdf")
    code = comments_main(
        [
            str(dump),
            "--config-dir",
            str(config_dir),
            "--output",
            str(dump / "report.xlsx"),
            "--no-pause",
        ]
    )
    assert code == 1


def test_default_report_name_uses_today(tmp_path: Path, config_dir: Path):
    dump = tmp_path / "comments"
    write_mbs_right_pdf(
        dump / "R459-MBS-ZZ-ZZ-DR-W-70001.pdf",
        document_reference="R459-MBS-ZZ-ZZ-DR-W-70001",
        title="Default report",
        revision="P01",
    )
    result = organise_comment_pdfs(dump, load_config(config_dir))
    path = write_comments_report(result)
    expected = dump / f"Comments_Organiser_{datetime.now().strftime('%d%m%y')}.xlsx"
    assert path == expected
    assert path.is_file()
